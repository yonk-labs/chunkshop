from __future__ import annotations

from pathlib import Path

from chunkshop.sources.parsers.base import ParserError
from chunkshop.sources.parsers.result import ParseResult


class DOCXParser:
    supported_extensions = ["docx"]

    def parse(self, path: Path) -> str | ParseResult:
        """Extract every top-level paragraph's text.

        The F5 claim here is deliberately NOT "the body has paragraphs".
        python-docx hands a brand-new blank document a single empty
        paragraph, so paragraph count would fire on a legitimately blank
        ``.docx`` — a false positive of exactly the kind this whole change
        exists to avoid. The honest signal is TABLES: this walk reads
        ``document.paragraphs`` and nothing else, so a document whose text
        all lives inside tables is content we demonstrably had and
        demonstrably failed to extract.
        """
        try:
            import docx
        except (ImportError, TypeError) as exc:
            raise RuntimeError(
                "DOCX parsing requires python-docx. Install with `pip install chunkshop[docx]`."
            ) from exc
        try:
            d = docx.Document(str(path))
            text = "\n".join(p.text for p in d.paragraphs)
            table_count = len(d.tables)
        except Exception as exc:
            raise ParserError(f"failed to parse DOCX {path}: {exc}") from exc
        if not text.strip() and table_count:
            return ParseResult(
                text=text, had_content=f"{table_count} table(s)"
            )
        return text
